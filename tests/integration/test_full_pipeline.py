from __future__ import annotations

import base64
import importlib.util
import tempfile
import unittest
from pathlib import Path


OPTIONAL_DEPS = {
    "docx": "python-docx",
    "numpy": "numpy",
    "PIL": "Pillow",
    "pydantic": "pydantic",
    "pypdfium2": "pypdfium2",
    "reportlab": "reportlab",
}


def _missing_optional_deps() -> list[str]:
    return [
        package_name
        for module_name, package_name in OPTIONAL_DEPS.items()
        if importlib.util.find_spec(module_name) is None
    ]


MISSING_OPTIONAL_DEPS = _missing_optional_deps()


@unittest.skipIf(
    MISSING_OPTIONAL_DEPS,
    "missing optional pipeline dependencies: " + ", ".join(MISSING_OPTIONAL_DEPS),
)
class FullPipelineSmokeTests(unittest.TestCase):
    def test_converting_outputs_chunk_with_valid_offsets(self):
        from PIL import Image, ImageDraw, ImageFont
        from docx import Document as DocxDocument
        from docx.shared import Inches
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.units import inch
        from reportlab.pdfgen import canvas
        from reportlab.platypus import Table, TableStyle

        from tiny_rag.chunking import SplitterConfig, split_with_diagnostics
        from tiny_rag.converting.parser.docx_parser import DocxParser
        from tiny_rag.converting.parser.markdown_parser import MarkdownParser
        from tiny_rag.converting.parser.pdf_parser import PDFParser

        with tempfile.TemporaryDirectory(prefix="tiny-rag-pipeline-") as tmp:
            raw_dir = Path(tmp) / "raw"
            raw_dir.mkdir()

            shared_image = raw_dir / "source_image.png"
            scanned_page = raw_dir / "scanned_page.png"
            _make_png(shared_image, "Shared Test Image", Image, ImageDraw, ImageFont)
            _make_png(scanned_page, "Scanned PDF Page", Image, ImageDraw, ImageFont)

            paths = [
                _make_markdown(raw_dir / "sample.md", shared_image),
                _make_txt(raw_dir / "sample.txt"),
                _make_docx(raw_dir / "sample.docx", shared_image, DocxDocument, Inches),
                _make_pdf(
                    raw_dir / "sample.pdf",
                    scanned_page,
                    canvas,
                    letter,
                    inch,
                    colors,
                    Table,
                    TableStyle,
                ),
            ]

            for path in paths:
                parser = _parser_for(path, MarkdownParser, DocxParser, PDFParser)
                document = parser.parse_into_text(path.read_bytes())
                self.assertTrue(document.content.strip(), path.name)

                chunks, diagnostics = split_with_diagnostics(
                    document.content,
                    SplitterConfig(
                        chunk_size=180,
                        chunk_overlap=30,
                        separators=["\n\n", "\n", ". ", "。"],
                        strategy="auto",
                    ),
                )

                self.assertTrue(chunks, path.name)
                self.assertIn(diagnostics.selected_tier, {"heading", "heuristic", "legacy"})
                self.assertEqual([], _offset_errors(document.content, chunks), path.name)

    def test_markdown_document_persists_chunks_and_indexes_for_retrieval(self):
        from tiny_rag.chunking import SplitterConfig
        from tiny_rag.converting.parser.markdown_parser import MarkdownParser
        from tiny_rag.indexing import (
            KEYWORDS_RETRIEVER_TYPE,
            VECTOR_RETRIEVER_TYPE,
            index_knowledge_after_chunks_persisted,
        )
        from tiny_rag.indexing.models import RetrieveParams
        from tiny_rag.indexing.repositories.sqlite import SQLiteIndexRepository
        from tiny_rag.persistence import ChunkRepository, connect, persist_text_chunks

        tenant_id = 7
        knowledge_id = "knowledge-e2e"
        knowledge_base_id = "kb-e2e"
        source = b"""# Retrieval Operations Runbook

## Ingest

The document parser should normalize tables before chunking begins.

| Metric | Value |
| --- | --- |
| latency_budget | 120ms |
| index_path | sqlite-vector |

## Search

Operators can retrieve the latency_budget term through keyword search after indexing.
Vector search should also return the stored chunk metadata.
"""

        document = MarkdownParser(file_name="runbook.md", file_type="md").parse_into_text(source)
        conn = connect(":memory:")
        persist_result = persist_text_chunks(
            conn,
            tenant_id=tenant_id,
            knowledge_id=knowledge_id,
            knowledge_base_id=knowledge_base_id,
            text=document.content,
            config=SplitterConfig(
                chunk_size=180,
                chunk_overlap=30,
                separators=["\n\n", "\n", ". "],
                strategy="auto",
            ),
        )

        repository = SQLiteIndexRepository(conn)
        embedder = _DeterministicEmbedder()
        index_stats = index_knowledge_after_chunks_persisted(
            conn=conn,
            repository=repository,
            embedder=embedder,
            knowledge_id=knowledge_id,
            title="Retrieval Operations Runbook",
            chunks=persist_result.inserted_chunks,
            retriever_types=[KEYWORDS_RETRIEVER_TYPE, VECTOR_RETRIEVER_TYPE],
        )

        persisted_text_chunks = ChunkRepository(conn).list_chunks_by_knowledge_id(
            tenant_id=tenant_id,
            knowledge_id=knowledge_id,
        )
        self.assertEqual(index_stats.text_chunk_count, len(persisted_text_chunks))
        self.assertEqual(index_stats.metadata_count, index_stats.text_chunk_count)
        self.assertEqual(index_stats.fts_count, index_stats.text_chunk_count)
        self.assertEqual(index_stats.vector_count, index_stats.text_chunk_count)
        self.assertEqual(index_stats.embedded_count, index_stats.text_chunk_count)

        keyword_results = repository.retrieve(
            RetrieveParams(
                query="latency_budget",
                top_k=3,
                knowledge_ids=[knowledge_id],
                retriever_type=KEYWORDS_RETRIEVER_TYPE,
            )
        )
        self.assertTrue(keyword_results[0].results)
        self.assertEqual(keyword_results[0].results[0].knowledge_id, knowledge_id)

        vector_results = repository.retrieve(
            RetrieveParams(
                embedding=embedder.embeddings[0],
                top_k=3,
                knowledge_ids=[knowledge_id],
                retriever_type=VECTOR_RETRIEVER_TYPE,
            )
        )
        text_chunk_ids = {chunk.id for chunk in persisted_text_chunks}
        self.assertTrue(vector_results[0].results)
        self.assertIn(vector_results[0].results[0].chunk_id, text_chunk_ids)


def _make_png(path: Path, title: str, Image, ImageDraw, ImageFont) -> None:
    img = Image.new("RGB", (640, 360), "white")
    draw = ImageDraw.Draw(img)
    try:
        font_title = ImageFont.truetype(
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
            34,
        )
        font_body = ImageFont.truetype(
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            22,
        )
    except Exception:
        font_title = ImageFont.load_default()
        font_body = ImageFont.load_default()
    draw.rectangle((24, 24, 616, 336), outline=(30, 40, 50), width=3)
    draw.text((60, 70), title, fill=(20, 30, 40), font=font_title)
    draw.text((60, 140), "Image extracted through converting.", fill=(20, 30, 40), font=font_body)
    img.save(path)


def _make_markdown(path: Path, image_path: Path) -> Path:
    payload = base64.b64encode(image_path.read_bytes()).decode()
    text = f"""# Markdown Manual

Intro paragraph for markdown conversion and chunk routing.

## Setup

- Install the parser package
- Run the converting stage
- Pass cleaned text into chunking

|Metric|Value|
|---|---|
|Formats|pdf, md, txt, docx|
|Images|inline base64|

![inline](data:image/png;base64,{payload})

## Verify

The final chunks should preserve offsets and keep table syntax readable.
"""
    path.write_text(text, encoding="utf-8")
    return path


def _make_txt(path: Path) -> Path:
    path.write_text(
        "\n\n".join(
            [
                "Plain Text Runbook",
                "Section 1 Overview\nThis text file is parsed through the markdown parser.",
                "Section 2 Checks\nOffsets, sequence numbers, and chunk content should line up.",
                "Section 3 Result\nThe final result should be ready for embedding.",
            ]
        ),
        encoding="utf-8",
    )
    return path


def _make_docx(path: Path, image_path: Path, DocxDocument, Inches) -> Path:
    doc = DocxDocument()
    doc.add_heading("Word Manual", level=1)
    doc.add_paragraph("This paragraph comes from a docx source document.")
    doc.add_heading("Setup", level=2)
    doc.add_paragraph("The parser should preserve useful text and table content.")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Format"
    table.cell(1, 1).text = "docx"
    table.cell(2, 0).text = "Stage"
    table.cell(2, 1).text = "converting"
    doc.add_picture(str(image_path), width=Inches(3.2))
    doc.add_heading("Verify", level=2)
    doc.add_paragraph("Chunking should run after the markdown content is produced.")
    doc.save(path)
    return path


def _make_pdf(path: Path, scan_image_path: Path, canvas, letter, inch, colors, Table, TableStyle) -> Path:
    c = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    margin = 54

    c.setFont("Helvetica-Bold", 24)
    c.drawString(margin, height - margin, "PDF Manual")
    c.setFont("Helvetica-Bold", 16)
    c.drawString(margin, height - margin - 42, "Executive Summary")
    y = height - margin - 70
    y = _draw_wrapped(
        c,
        "The PDF parser should extract native text, flatten table content, and route image-only pages to image references.",
        margin,
        y,
        width - 2 * margin,
    )
    y -= 10
    for item in ["Extract native text", "Detect scanned pages", "Keep image references"]:
        c.setFont("Helvetica", 10.5)
        c.drawString(margin + 12, y, "- " + item)
        y -= 16
    y -= 8

    table = Table(
        [["Metric", "Value"], ["Format", "pdf"], ["Scanned pages", "1"]],
        colWidths=[2.0 * inch, 2.0 * inch],
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF7")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#7A8795")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    _, table_h = table.wrapOn(c, width - 2 * margin, height)
    table.drawOn(c, margin, y - table_h)
    c.showPage()

    c.drawImage(
        str(scan_image_path),
        36,
        36,
        width=width - 72,
        height=height - 72,
        preserveAspectRatio=True,
        anchor="c",
    )
    c.save()
    return path


def _draw_wrapped(c, text: str, x: float, y: float, max_width: float) -> float:
    font = "Helvetica"
    size = 10.5
    leading = 14
    c.setFont(font, size)
    words = text.split()
    line = ""
    for word in words:
        probe = word if not line else line + " " + word
        if c.stringWidth(probe, font, size) <= max_width:
            line = probe
        else:
            c.drawString(x, y, line)
            y -= leading
            line = word
    if line:
        c.drawString(x, y, line)
        y -= leading
    return y


def _parser_for(path: Path, MarkdownParser, DocxParser, PDFParser):
    ext = path.suffix.lower().lstrip(".")
    if ext == "pdf":
        return PDFParser(file_name=path.name, file_type=ext)
    if ext == "docx":
        return DocxParser(file_name=path.name, file_type=ext)
    if ext in {"md", "markdown", "txt"}:
        return MarkdownParser(file_name=path.name, file_type=ext)
    raise ValueError(ext)


def _offset_errors(text: str, chunks) -> list[str]:
    errors: list[str] = []
    for chunk in chunks:
        if text[chunk.start : chunk.end] != chunk.content:
            errors.append(f"seq={chunk.seq} offset mismatch [{chunk.start}:{chunk.end}]")
        if chunk.end - chunk.start != len(chunk.content):
            errors.append(f"seq={chunk.seq} length mismatch")
    return errors


class _DeterministicEmbedder:
    def __init__(self, dimensions: int = 4):
        self.dimensions = dimensions
        self.batch_calls: list[list[str]] = []
        self.embeddings: list[list[float]] = []

    def batch_embed(self, texts: list[str]) -> list[list[float]]:
        self.batch_calls.append(list(texts))
        embeddings = [_vector_for_text(text, self.dimensions) for text in texts]
        self.embeddings.extend(embeddings)
        return embeddings

    def get_model_name(self) -> str:
        return "deterministic-test"

    def get_dimensions(self) -> int:
        return self.dimensions

    def get_model_id(self) -> str:
        return "deterministic-test"


def _vector_for_text(text: str, dimensions: int) -> list[float]:
    seed = sum((index + 1) * ord(char) for index, char in enumerate(text[:512]))
    features = [
        float(seed % 97 + 1),
        float(len(text) % 89 + 1),
        float(text.count("\n") + 1),
        float(len(set(text)) % 83 + 1),
    ]
    if dimensions <= len(features):
        return features[:dimensions]
    return features + [1.0] * (dimensions - len(features))


if __name__ == "__main__":
    unittest.main()
